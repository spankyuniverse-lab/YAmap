"""Экспорт данных: CSV, JSON, XLSX."""

from __future__ import annotations

import csv
import json
import logging
import re
from dataclasses import asdict
from pathlib import Path

from .models import FIELDNAMES, HEADERS_RU, Organization

log = logging.getLogger("yandex_parser")


def save_csv(orgs: list[Organization], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES, delimiter=";")
        writer.writeheader()
        for org in orgs:
            writer.writerow(asdict(org))
    log.info("CSV сохранён: %s (%d записей)", path, len(orgs))


def save_json(orgs: list[Organization], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    data = {
        "total": len(orgs),
        "organizations": [asdict(org) for org in orgs],
    }
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    log.info("JSON сохранён: %s (%d записей)", path, len(orgs))


def _write_sheet(ws, orgs: list[Organization]) -> None:
    """Записать организации на один лист Excel."""
    from openpyxl.styles import Font

    cols = FIELDNAMES

    for col_idx, field_name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=HEADERS_RU.get(field_name, field_name))
        cell.font = Font(bold=True)

    for row_idx, org in enumerate(orgs, 2):
        d = asdict(org)
        for col_idx, field_name in enumerate(cols, 1):
            ws.cell(row=row_idx, column=col_idx, value=d.get(field_name, ""))

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)


def save_xlsx(orgs: list[Organization], path: Path) -> None:
    try:
        from openpyxl import Workbook
    except ImportError:
        log.warning("openpyxl не установлен — сохраняю в CSV")
        save_csv(orgs, path.with_suffix(".csv"))
        return

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Яндекс.Карты"
    _write_sheet(ws, orgs)
    wb.save(path)
    log.info("XLSX сохранён: %s (%d записей)", path, len(orgs))


def save_xlsx_by_categories(
    results: dict[str, list[Organization]],
    path: Path,
) -> None:
    """Сохранить результаты по категориям: отдельный лист на каждую + сводный."""
    try:
        from openpyxl import Workbook
    except ImportError:
        log.warning("openpyxl не установлен — сохраняю сводный CSV")
        all_orgs = []
        for orgs in results.values():
            all_orgs.extend(orgs)
        save_csv(all_orgs, path.with_suffix(".csv"))
        return

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws_all = wb.active
    ws_all.title = "Все результаты"
    all_orgs: list[Organization] = []
    for orgs in results.values():
        all_orgs.extend(orgs)
    _write_sheet(ws_all, all_orgs)

    for query, orgs in results.items():
        if not orgs:
            continue
        sheet_name = re.sub(r'[\\/*?\[\]:]', '', query)[:31]
        ws = wb.create_sheet(title=sheet_name)
        _write_sheet(ws, orgs)

    wb.save(path)
    total = sum(len(v) for v in results.values())
    log.info(
        "XLSX сохранён: %s (%d записей, %d листов)",
        path, total, len(wb.sheetnames),
    )


def save_auto(orgs: list[Organization], out_path: Path) -> None:
    """Сохранить в формат по расширению файла."""
    suffix = out_path.suffix.lower()
    if suffix == ".csv":
        save_csv(orgs, out_path)
    elif suffix == ".json":
        save_json(orgs, out_path)
    else:
        save_xlsx(orgs, out_path)
