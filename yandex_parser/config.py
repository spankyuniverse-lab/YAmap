"""Конфигурация, категории, ResumeManager."""

from __future__ import annotations

import csv
import json
import logging
from pathlib import Path
from typing import Any

from .models import (
    FIELDNAMES,
    HEADERS_RU,
    Organization,
    _normalize_for_dedup,
)

log = logging.getLogger("yandex_parser")


# ---------------------------------------------------------------------------
# Config file support (YAML/JSON)
# ---------------------------------------------------------------------------

DEFAULT_CONFIG: dict[str, Any] = {
    "max_results": 500,
    "scroll_pause": 1.0,
    "headless": True,
    "detail": False,
    "api_intercept": False,
    "output": "results.xlsx",
    "proxy": None,
    "proxy_file": None,
    "city": None,
    "categories": None,
}


def load_config(path: str) -> dict[str, Any]:
    """Загрузить конфиг из YAML или JSON файла."""
    p = Path(path)
    if not p.exists():
        log.warning("Конфиг-файл не найден: %s", path)
        return {}

    text = p.read_text(encoding="utf-8")
    if p.suffix in (".yaml", ".yml"):
        try:
            import yaml
            return yaml.safe_load(text) or {}
        except ImportError:
            log.warning("PyYAML не установлен — пробую как JSON")
    try:
        return json.loads(text)
    except Exception as exc:
        log.warning("Ошибка чтения конфига %s: %s", path, exc)
        return {}


# ---------------------------------------------------------------------------
# Каталог категорий (аналог 2ГИС)
# ---------------------------------------------------------------------------

CATEGORIES: dict[str, list[str]] = {
    "еда": [
        "рестораны", "кафе", "бары", "пиццерии", "суши-бары",
        "столовые", "фастфуд", "кофейни", "кондитерские", "пекарни",
        "шаурма", "бургерные", "доставка еды",
    ],
    "продукты": [
        "продуктовые магазины", "супермаркеты", "мясные магазины",
        "рыбные магазины", "овощи и фрукты", "молочные продукты",
        "алкогольные магазины", "кулинарии",
    ],
    "здоровье": [
        "аптеки", "больницы", "поликлиники", "стоматологии",
        "медицинские центры", "ветеринарные клиники", "лаборатории",
        "оптика", "косметология",
    ],
    "авто": [
        "автосервисы", "шиномонтаж", "автомойки", "автозапчасти",
        "АЗС", "автосалоны", "эвакуаторы", "парковки",
        "техосмотр", "автострахование",
    ],
    "красота": [
        "салоны красоты", "парикмахерские", "барбершопы",
        "маникюр", "массаж", "спа-салоны", "солярии", "тату-салоны",
    ],
    "покупки": [
        "торговые центры", "магазины одежды", "магазины обуви",
        "магазины электроники", "магазины мебели", "строительные магазины",
        "магазины цветов", "зоомагазины", "книжные магазины",
        "магазины подарков", "ювелирные магазины",
    ],
    "услуги": [
        "банки", "банкоматы", "нотариусы", "юридические услуги",
        "страховые компании", "фотостудии", "ателье", "химчистки",
        "ремонт телефонов", "ремонт бытовой техники", "клининг",
        "курьерские службы", "типографии",
    ],
    "образование": [
        "школы", "детские сады", "университеты", "колледжи",
        "языковые курсы", "автошколы", "репетиторы",
        "курсы программирования", "музыкальные школы",
    ],
    "спорт": [
        "фитнес-клубы", "тренажёрные залы", "бассейны",
        "спортивные магазины", "йога-студии", "танцевальные студии",
        "боксёрские клубы", "теннисные корты", "спортивные площадки",
    ],
    "развлечения": [
        "кинотеатры", "театры", "музеи", "парки развлечений",
        "боулинг", "бильярд", "караоке", "квесты",
        "ночные клубы", "концертные залы",
    ],
    "туризм": [
        "гостиницы", "хостелы", "турагентства", "достопримечательности",
        "экскурсии", "аренда автомобилей", "визовые центры",
    ],
    "транспорт": [
        "такси", "каршеринг", "автобусные станции",
        "железнодорожные вокзалы", "аэропорты", "грузоперевозки",
    ],
    "недвижимость": [
        "агентства недвижимости", "новостройки",
        "управляющие компании", "жилые комплексы",
    ],
    "дом": [
        "мебельные магазины", "сантехника", "электрика",
        "окна и двери", "кухни на заказ", "натяжные потолки",
        "кондиционеры", "отопление",
    ],
    "IT": [
        "компьютерные магазины", "ремонт компьютеров",
        "IT-компании", "интернет-провайдеры", "веб-студии",
    ],
}


def list_categories() -> None:
    """Вывести каталог категорий в консоль."""
    print("\n  Каталог категорий (аналог 2ГИС):\n")
    for group, items in CATEGORIES.items():
        print(f"  [{group}]")
        for item in items:
            print(f"    - {item}")
        print()
    print("Использование:")
    print('  python -m yandex_parser --city Москва --category еда')
    print('  python -m yandex_parser --city Москва --category еда рестораны кафе')
    print('  python -m yandex_parser --city Москва --all-categories')


def resolve_categories(names: list[str]) -> list[str]:
    """Преобразовать названия групп/категорий в список поисковых запросов."""
    queries: list[str] = []
    for name in names:
        key = name.lower().strip()
        if key in CATEGORIES:
            queries.extend(CATEGORIES[key])
        else:
            queries.append(name.strip())
    return queries


# ---------------------------------------------------------------------------
# Resume manager
# ---------------------------------------------------------------------------

class ResumeManager:
    """Менеджер докачки: загружает уже собранные данные из файла."""

    def __init__(self, path: Path | None = None):
        self._existing: dict[str, Organization] = {}
        self._completed_queries: set[str] = set()
        if path and path.exists():
            self._load(path)

    def _load(self, path: Path) -> None:
        suffix = path.suffix.lower()
        rows: list[dict] = []

        if suffix == ".csv":
            try:
                with open(path, "r", encoding="utf-8-sig") as f:
                    reader = csv.DictReader(f, delimiter=";")
                    rows = list(reader)
            except Exception as exc:
                log.warning("Не удалось прочитать CSV для резюме: %s", exc)
                return

        elif suffix == ".xlsx":
            try:
                from openpyxl import load_workbook
                wb = load_workbook(path, read_only=True)
                ws = wb.active
                if ws is None:
                    return
                headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    row_dict = {h: (v or "") for h, v in zip(headers, row) if h}
                    rows.append(row_dict)
                wb.close()
            except Exception as exc:
                log.warning("Не удалось прочитать XLSX для резюме: %s", exc)
                return

        elif suffix == ".json":
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                if isinstance(data, list):
                    rows = data
                elif isinstance(data, dict) and "organizations" in data:
                    rows = data["organizations"]
            except Exception as exc:
                log.warning("Не удалось прочитать JSON для резюме: %s", exc)
                return

        ru_to_en = {v: k for k, v in HEADERS_RU.items()}

        for row in rows:
            norm = {}
            for k, v in row.items():
                en_key = ru_to_en.get(k, k)
                norm[en_key] = str(v) if v else ""

            key = (
                f"{_normalize_for_dedup(norm.get('name', ''))}"
                f"|{_normalize_for_dedup(norm.get('address', ''))}"
            )
            if key != "|":
                org = Organization(**{f: norm.get(f, "") for f in FIELDNAMES if f in norm})
                self._existing[key] = org
                q = norm.get("search_query", "")
                if q:
                    self._completed_queries.add(q)

        log.info(
            "Резюме: загружено %d существующих организаций, %d выполненных запросов",
            len(self._existing), len(self._completed_queries),
        )

    @property
    def existing_count(self) -> int:
        return len(self._existing)

    def is_known(self, name: str, address: str) -> bool:
        return (
            f"{_normalize_for_dedup(name)}|{_normalize_for_dedup(address)}"
            in self._existing
        )

    def is_query_done(self, query: str) -> bool:
        return query in self._completed_queries

    def existing_orgs(self) -> list[Organization]:
        return list(self._existing.values())

    def existing_keys(self) -> set[str]:
        return set(self._existing.keys())
