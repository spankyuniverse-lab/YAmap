# -*- coding: utf-8 -*-
"""Живой прогон опций --fast-api и --url-only против макета Яндекса.

Главное, что проверяем: обе опции дают ТОТ ЖЕ результат, что обычный сбор.
Ускоритель, который теряет организации, хуже медленного сбора.
"""
import os
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import mock_yandex

CATALOG = {"Заправки": 20, "Кафе": 8}
fails = []


def check(cond, msg, detail=""):
    print(("OK   | " if cond else "FAIL | ") + msg + (f"\n       {detail}" if not cond and detail else ""))
    if not cond:
        fails.append(msg)


def rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Все результаты"] if "Все результаты" in wb.sheetnames else wb.active
    head = [c.value for c in ws[1]]
    return [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)
            if any(c.value for c in r)]


def run(args, env):
    return subprocess.run(
        [sys.executable, str(HERE.parent / "yandex_parser.py"), *args],
        cwd=HERE, env=env, capture_output=True, text=True, timeout=600)


def main() -> int:
    browser = os.environ.get("YAMAP_BROWSER_PATH")
    if not browser:
        for cand in ("/opt/pw-browsers/chromium",
                     "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"):
            if Path(cand).exists():
                browser = cand
                break
    if not browser:
        print("ПРОПУСК: браузер не найден (playwright/chromium не установлен)")
        return 0

    base, srv = mock_yandex.start(CATALOG)
    print(f"макет Яндекса поднят на {base}\n")
    env = dict(os.environ)
    env.update({"YAMAP_BASE_URL": base, "YAMAP_BROWSER_PATH": browser,
                "YAMAP_HEADLESS": "1", "PYTHONIOENCODING": "utf-8"})

    for junk in HERE.glob("fa_*.xlsx"):
        junk.unlink()

    try:
        base_args = ["--city", "Алматы", "--category", "Заправки", "Кафе", "-n", "500"]

        # ---- 1. Эталон: обычный сбор ----
        print("--- обычный сбор (эталон) ---")
        out0 = HERE / "fa_plain.xlsx"
        r = run([*base_args, "-o", str(out0), "--api-intercept"], env)
        check(r.returncode == 0, f"обычный сбор отработал (код {r.returncode})",
              (r.stdout + r.stderr)[-600:])
        plain = rows(out0) if out0.exists() else []
        check(len(plain) == 28, f"эталон: собрано {len(plain)}, ждали 28 (20+8)")
        plain_names = {r.get("Название") for r in plain}

        # ---- 2. --url-only ----
        print("\n--- --url-only ---")
        out1 = HERE / "fa_urlonly.xlsx"
        r = run([*base_args, "-o", str(out1), "--api-intercept", "--url-only"], env)
        check(r.returncode == 0, f"--url-only отработал (код {r.returncode})",
              (r.stdout + r.stderr)[-600:])
        uo = rows(out1) if out1.exists() else []
        check(len(uo) == len(plain), f"--url-only: {len(uo)} строк, как в эталоне {len(plain)}")
        check({r.get("Название") for r in uo} == plain_names,
              "--url-only: тот же состав организаций")
        check("иду по прямому URL" in (r.stdout + r.stderr)
              or "url-only" in (r.stdout + r.stderr).lower(),
              "--url-only: режим виден в логе")

        # ---- 3. --fast-api ----
        print("\n--- --fast-api ---")
        out2 = HERE / "fa_api.xlsx"
        r = run([*base_args, "-o", str(out2), "--fast-api"], env)
        log = r.stdout + r.stderr
        check(r.returncode == 0, f"--fast-api отработал (код {r.returncode})", log[-600:])
        fa = rows(out2) if out2.exists() else []
        check("fast-api:" in log, "--fast-api: быстрый путь реально сработал",
              log[-400:])
        check(len(fa) == len(plain), f"--fast-api: {len(fa)} строк, как в эталоне {len(plain)}")
        check({r.get("Название") for r in fa} == plain_names,
              "--fast-api: тот же состав организаций")
        check(all(r.get("Адрес") for r in fa), "--fast-api: адреса на месте")
        check(sum(1 for r in fa if r.get("Категория")) >= len(fa) * 0.9,
              "--fast-api: категории на месте")

        # ---- 4. Пагинация: 20 заправок не влезают в одну страницу (12) ----
        check("стр." in log, "--fast-api: выдача снята постранично", log[-300:])

        # ---- 5. Откат: API недоступен → обычный сбор, ничего сверх не теряем.
        # Гасим API у макета. Важно: у макета от того же эндпоинта зависит и
        # его собственная кнопка «Показать ещё», так что обычный сбор в этих
        # условиях тоже недосчитается страниц. Поэтому эталон для отката —
        # прогон БЕЗ --fast-api при том же погашенном API: --fast-api обязан
        # дать ровно столько же, то есть не потерять ничего сверх обстановки.
        print("\n--- --fast-api при недоступном API (откат) ---")
        mock_yandex.DISABLE_API = True
        try:
            out4 = HERE / "fa_broken_plain.xlsx"
            r0 = run([*base_args, "-o", str(out4), "--api-intercept"], env)
            broken_plain = rows(out4) if out4.exists() else []
            check(r0.returncode == 0, "эталон при погашенном API отработал")

            out3 = HERE / "fa_fallback.xlsx"
            r = run([*base_args, "-o", str(out3), "--fast-api", "--api-intercept"], env)
            fb = rows(out3) if out3.exists() else []
        finally:
            mock_yandex.DISABLE_API = False

        check(r.returncode == 0, f"откат отработал без ошибки (код {r.returncode})")
        check(len(broken_plain) > 0, f"эталон при погашенном API непустой ({len(broken_plain)})")
        check(len(fb) == len(broken_plain),
              f"откат: собрано {len(fb)}, как без --fast-api в тех же условиях "
              f"({len(broken_plain)}) — быстрый путь ничего не потерял")
        check({r.get("Название") for r in fb} == {r.get("Название") for r in broken_plain},
              "откат: тот же состав организаций, что и без --fast-api")
        check("обычным способом" in (r.stdout + r.stderr),
              "откат: парсер сам сообщил о переходе на обычный сбор")

    finally:
        srv.shutdown()
        for junk in HERE.glob("fa_*.xlsx"):
            junk.unlink()

    print("\n" + ("✅ УСКОРИТЕЛИ ОК" if not fails else f"❌ ПРОВАЛЫ: {fails}"))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
