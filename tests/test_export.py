"""Тесты: экспорт CSV, JSON, XLSX."""

import csv
import json
import pytest
from pathlib import Path

from yandex_parser.models import Organization
from yandex_parser.export import save_auto, save_csv, save_json, save_xlsx


def _sample_orgs() -> list[Organization]:
    return [
        Organization(
            name="Кафе Рога",
            address="ул. Ленина, 5",
            phone="+7 (495) 123-45-67",
            rating="4.5",
            category="кафе",
        ),
        Organization(
            name="Бар Луна",
            address="пр-т Мира, 10",
            website="luna-bar.ru",
            reviews_count="42",
        ),
    ]


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


class TestSaveCSV:
    def test_creates_file(self, tmp_path):
        path = tmp_path / "out.csv"
        save_csv(_sample_orgs(), path)
        assert path.exists()

    def test_content(self, tmp_path):
        path = tmp_path / "out.csv"
        save_csv(_sample_orgs(), path)

        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f, delimiter=";")
            rows = list(reader)

        assert len(rows) == 2
        assert rows[0]["name"] == "Кафе Рога"
        assert rows[0]["phone"] == "+7 (495) 123-45-67"
        assert rows[1]["name"] == "Бар Луна"

    def test_empty(self, tmp_path):
        path = tmp_path / "empty.csv"
        save_csv([], path)
        assert path.exists()
        with open(path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f, delimiter=";")
            rows = list(reader)
        assert len(rows) == 0

    def test_creates_parent_dirs(self, tmp_path):
        path = tmp_path / "sub" / "dir" / "out.csv"
        save_csv(_sample_orgs(), path)
        assert path.exists()


# ---------------------------------------------------------------------------
# JSON
# ---------------------------------------------------------------------------


class TestSaveJSON:
    def test_creates_file(self, tmp_path):
        path = tmp_path / "out.json"
        save_json(_sample_orgs(), path)
        assert path.exists()

    def test_content(self, tmp_path):
        path = tmp_path / "out.json"
        save_json(_sample_orgs(), path)

        data = json.loads(path.read_text(encoding="utf-8"))
        assert data["total"] == 2
        assert len(data["organizations"]) == 2
        assert data["organizations"][0]["name"] == "Кафе Рога"

    def test_empty(self, tmp_path):
        path = tmp_path / "empty.json"
        save_json([], path)
        data = json.loads(path.read_text(encoding="utf-8"))
        assert data["total"] == 0
        assert data["organizations"] == []

    def test_utf8(self, tmp_path):
        path = tmp_path / "utf.json"
        save_json(_sample_orgs(), path)
        text = path.read_text(encoding="utf-8")
        assert "Кафе Рога" in text  # ensure_ascii=False


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


class TestSaveXLSX:
    def test_creates_file(self, tmp_path):
        path = tmp_path / "out.xlsx"
        save_xlsx(_sample_orgs(), path)
        assert path.exists()

    def test_content(self, tmp_path):
        from openpyxl import load_workbook

        path = tmp_path / "out.xlsx"
        save_xlsx(_sample_orgs(), path)

        wb = load_workbook(path)
        ws = wb.active
        # Header row
        headers = [cell.value for cell in ws[1]]
        assert "Название" in headers
        assert "Телефон" in headers

        # Data
        assert ws.cell(row=2, column=1).value == "Кафе Рога"
        assert ws.cell(row=3, column=1).value == "Бар Луна"
        wb.close()


# ---------------------------------------------------------------------------
# save_auto
# ---------------------------------------------------------------------------


class TestSaveAuto:
    def test_csv(self, tmp_path):
        path = tmp_path / "auto.csv"
        save_auto(_sample_orgs(), path)
        assert path.exists()

    def test_json(self, tmp_path):
        path = tmp_path / "auto.json"
        save_auto(_sample_orgs(), path)
        data = json.loads(path.read_text(encoding="utf-8"))
        assert data["total"] == 2

    def test_xlsx(self, tmp_path):
        path = tmp_path / "auto.xlsx"
        save_auto(_sample_orgs(), path)
        assert path.exists()
