# -*- coding: utf-8 -*-
"""Ускорители не должны терять данные.

Быстрее — легко: перестань ждать. Вопрос в том, не обрывается ли при этом
выдача. Гоняем парсер по макету на размерах ВОКРУГ границы страницы (12) и
сверяем с эталоном, который знает макет.
"""
import os
import re
import subprocess
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import mock_yandex

# Размеры взяты вокруг PAGE_SIZE=12: ровно страница, страница+1, две
# страницы, две+1 — там и обрывается сбор, если поспешить с выходом.
CATALOG = {"Нольовое": 0, "Одиночное": 1, "Ровно страница": 12,
           "Страница плюс один": 13, "Две страницы": 24,
           "Две плюс один": 25}
fails = []


def check(cond, msg, detail=""):
    print(("OK   | " if cond else "FAIL | ") + msg
          + (f"\n       {detail}" if not cond and detail else ""))
    if not cond:
        fails.append(msg)


def main() -> int:
    browser = os.environ.get("YAMAP_BROWSER_PATH")
    if not browser:
        for cand in ("/opt/pw-browsers/chromium",
                     "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"):
            if Path(cand).exists():
                browser = cand
                break
    if not browser:
        print("ПРОПУСК: браузер не найден (playwright/chromium не установлен)")
        return 0

    base, srv = mock_yandex.start(CATALOG)
    print(f"макет Яндекса поднят на {base}\n")
    env = dict(os.environ)
    env.update({"YAMAP_BASE_URL": base, "YAMAP_BROWSER_PATH": browser,
                "YAMAP_HEADLESS": "1", "PYTHONIOENCODING": "utf-8"})
    out = HERE / "speed.xlsx"
    for junk in (out, HERE / "speed.cities.json"):
        if junk.exists():
            junk.unlink()

    try:
        t0 = time.time()
        r = subprocess.run(
            [sys.executable, str(HERE.parent / "yandex_parser.py"),
             "--city", "Алматы", "--category", *CATALOG.keys(),
             "-n", "500", "-o", str(out), "--api-intercept"],
            cwd=HERE, env=env, capture_output=True, text=True, timeout=900)
        elapsed = time.time() - t0
        log = r.stdout + r.stderr
        check(r.returncode == 0, f"прогон отработал (код {r.returncode})",
              log[-1500:])
        if not out.exists():
            check(False, "книга создана")
            return 1

        from openpyxl import load_workbook
        ws = load_workbook(out)["Все результаты"]
        head = [c.value for c in ws[1]]
        qcol = head.index("Поисковый запрос")
        got: dict[str, int] = {}
        for row in ws.iter_rows(min_row=2):
            q = row[qcol].value
            if q:
                got[q] = got.get(q, 0) + 1

        # --- ГЛАВНОЕ: ни одна выдача не обрезана ---
        for query, expected in CATALOG.items():
            check(got.get(query, 0) == expected,
                  f"«{query}»: собрано {got.get(query, 0)}, в макете {expected}")

        # --- ранний выход действительно сработал ---
        check("Выдача кончилась" in log,
              "короткие выдачи закрываются по признаку «грузить нечего», "
              "а не по десяти холостым кругам")

        # --- честный ноль не перезапрашивается трижды ---
        check("ничего не найдено" in log,
              "пустая рубрика опознана как честный ноль")
        check("попытка 1/3" not in log and "попытка 2/3" not in log,
              "честный ноль не гоняется тремя попытками",
              re.findall(r".*попытка \d/3.*", log)[:3])

        # --- лишний проход API не делается, когда API молчит ---
        check("Доп. проход API пропущен" in log or "API-обогащение" in log,
              "холостой доп. проход API пропускается")

        total = sum(CATALOG.values())
        print(f"\n   собрано {sum(got.values())} из {total} за {elapsed:.0f} c "
              f"({len(CATALOG)} рубрик)")

        # ---- браузер поднимается ОДИН раз на весь прогон, а не на каждый НП ----
        print("\n--- три НП подряд ---")
        multi = HERE / "speed_multi.xlsx"
        for junk in list(HERE.glob("speed_multi*")):
            junk.unlink() if junk.is_file() else None
        t1 = time.time()
        r2 = subprocess.run(
            [sys.executable, str(HERE.parent / "yandex_parser.py"),
             "--all-cities", "--cities", "Алматы,Астана,Шымкент",
             "--category", "Одиночное", "-n", "50", "-o", str(multi)],
            cwd=HERE, env=env, capture_output=True, text=True, timeout=900)
        el2 = time.time() - t1
        log2 = r2.stdout + r2.stderr
        check(r2.returncode == 0, f"прогон по трём НП отработал (код {r2.returncode})",
              log2[-1200:])
        warmups = log2.count("Прогрев: естественная навигация")
        check(warmups == 1,
              f"Chrome прогревается ОДИН раз на прогон, а не на каждый НП "
              f"(прогревов: {warmups}, НП: 3)")
        check("═══ Город 3/3" in log2, "все три НП пройдены")
        if multi.exists():
            ws2 = load_workbook(multi)["Все результаты"]
            check(ws2.max_row - 1 == 3,
                  f"по одной организации с каждого НП: {ws2.max_row - 1} из 3")
        print(f"   три НП за {el2:.0f} c")
        for junk in list(HERE.glob("speed_multi*")):
            if junk.is_file():
                junk.unlink()
        import shutil
        shutil.rmtree(HERE / "speed_multi_parts", ignore_errors=True)
    finally:
        srv.shutdown()
        for junk in (out, HERE / "speed.cities.json"):
            if junk.exists():
                junk.unlink()

    print("\n" + ("✅ УСКОРЕНИЕ БЕЗ ПОТЕРЬ" if not fails else f"❌ ПРОВАЛЫ: {fails}"))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
